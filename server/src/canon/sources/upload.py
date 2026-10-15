"""Files a person uploads: Markdown, text, HTML, PDF, Word (.docx), and zip archives of those.

Each file is an `upload` document dated by its modified time (the browser's `lastModified`, or the time inside the
zip), titled by its first heading or its name. Re-uploading a file of the same name replaces its passages, so the
facts the old version stated are closed rather than lost.

Safety, since the bytes come from anyone with a workspace:

- at most `MAX_FILES` files and `MAX_TOTAL` bytes an upload, `MAX_FILE` bytes a file;
- a zip is read in memory, never extracted to disk, so its paths cannot escape anywhere. Nested archives are
  skipped. Each entry's declared size, its compression ratio and the archive's total unpacked size are bounded,
  and entries are read through a capped stream, so a zip bomb stops at the cap whatever its headers claim;
- PDFs are read to `MAX_PDF_PAGES` pages; text is cut at `MAX_CHARS` characters a file.
"""

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import PurePosixPath

from canon.sources.client import SourceError
from canon.sources.connector import Report
from canon.sources.identity import SourceResult
from canon.sources.richtext import html_title, html_to_text
from canon.sources.text import clean_mdx, frontmatter, passages_of, slugify, split_passages
from canon.types import Batch, Doc

MAX_FILES = 200
MAX_FILE = 10 * 1024 * 1024
MAX_TOTAL = 30 * 1024 * 1024
MAX_ZIP_ENTRIES = 500
MAX_UNZIPPED = 60 * 1024 * 1024
MAX_RATIO = 100
MAX_PDF_PAGES = 300
MAX_CHARS = 400_000
TYPES = (".md", ".markdown", ".mdx", ".txt", ".html", ".htm", ".pdf", ".docx")


@dataclass(frozen=True)
class File:
    name: str
    data: bytes
    modified: str | None = None  # ISO time


def safe_name(name: str) -> str:
    """A file's path inside an upload, with no way out of it: no drive, no `..`, no leading slash."""
    parts = [p for p in PurePosixPath(name.replace("\\", "/")).parts if p not in ("", ".", "..", "/")]
    return "/".join(parts)[-200:] or "file"


def _iso(ts: float | None) -> str | None:
    if not ts:
        return None
    return datetime.fromtimestamp(ts, UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


def expand(files: list[File], notes: list[str]) -> list[File]:
    """Zips opened into their files, everything checked against the limits."""
    if len(files) > MAX_FILES:
        raise SourceError(f"Upload at most {MAX_FILES} files at a time.")
    if sum(len(f.data) for f in files) > MAX_TOTAL:
        raise SourceError(f"An upload can be at most {MAX_TOTAL // (1024 * 1024)} MB.")
    out: list[File] = []
    for f in files:
        name = safe_name(f.name)
        if len(f.data) > MAX_FILE and not name.lower().endswith(".zip"):
            notes.append(f"{name} is larger than {MAX_FILE // (1024 * 1024)} MB and was skipped.")
        elif name.lower().endswith(".zip"):
            out += _unzip(name, f.data, notes)
        elif name.lower().endswith(TYPES):
            out.append(File(name, f.data, f.modified))
        else:
            notes.append(f"{name} is not a file Canon reads (Markdown, text, HTML, PDF or Word) and was skipped.")
    if len(out) > MAX_FILES:
        notes.append(f"Read the first {MAX_FILES} files.")
        out = out[:MAX_FILES]
    return out


def _unzip(name: str, data: bytes, notes: list[str]) -> list[File]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except (zipfile.BadZipFile, ValueError):
        notes.append(f"{name} is not a zip archive Canon can open.")
        return []
    out: list[File] = []
    total = 0
    with archive:
        entries = [e for e in archive.infolist() if not e.is_dir()]
        if len(entries) > MAX_ZIP_ENTRIES:
            notes.append(f"{name} has {len(entries)} files; the first {MAX_ZIP_ENTRIES} were read.")
            entries = entries[:MAX_ZIP_ENTRIES]
        for e in entries:
            inner = safe_name(e.filename)
            if inner.split("/")[0] == "__MACOSX" or PurePosixPath(inner).name.startswith("."):
                continue
            if not inner.lower().endswith(TYPES):
                continue  # nested archives too
            if e.flag_bits & 0x1:
                notes.append(f"{name}: {inner} is encrypted and was skipped.")
                continue
            if e.file_size > MAX_FILE or (e.compress_size and e.file_size / e.compress_size > MAX_RATIO):
                notes.append(f"{name}: {inner} unpacks too large and was skipped.")
                continue
            if total + e.file_size > MAX_UNZIPPED:
                notes.append(f"{name} unpacks to more than {MAX_UNZIPPED // (1024 * 1024)} MB; the rest was skipped.")
                break
            try:
                with archive.open(e) as stream:
                    body = stream.read(MAX_FILE + 1)  # never more than the cap, whatever the header said
            except (zipfile.BadZipFile, RuntimeError, NotImplementedError, OSError):
                notes.append(f"{name}: {inner} could not be read.")
                continue
            if len(body) > MAX_FILE:
                notes.append(f"{name}: {inner} unpacks too large and was skipped.")
                continue
            total += len(body)
            try:
                modified = datetime(*e.date_time, tzinfo=UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
            except ValueError:
                modified = None
            stem = name.removesuffix(".zip")
            # `handbook.zip` holding `handbook/deploys.md` is `handbook/deploys.md`, not `handbook/handbook/...`.
            path = inner if inner.startswith(PurePosixPath(stem).name + "/") else f"{stem}/{inner}"
            out.append(File(path, body, modified))
    return out


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pdf(data: bytes) -> tuple[str, str | None]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("encrypted") from None
    title = None
    try:
        title = (reader.metadata or {}).get("/Title") or None
    except Exception:
        pass
    pages = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    text = "\n\n".join(pages)
    text = re.sub(r"-\n(?=[a-z])", "", text)  # words hyphenated across lines
    text = re.sub(r"(?<![.!?:\n])\n(?!\n)", " ", text)  # lines wrapped mid-sentence
    return text, str(title) if title else None


def _docx(data: bytes) -> tuple[str, str | None]:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = []
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        m = re.match(r"Heading (\d)", style)
        if style == "Title":
            lines.append(f"# {text}")
        elif m:
            lines.append("#" * min(6, int(m.group(1)) + 1) + " " + text)
        elif "List" in style:
            lines.append(f"- {text}")
        else:
            lines.append(f"\n{text}\n")
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(dict.fromkeys(cells)))
    title = (document.core_properties.title or "").strip() or None
    return "\n".join(lines), title


def text_of(f: File) -> tuple[str, str | None]:
    """A file's text (Markdown-like, with `#` headings) and its own title when it has one."""
    lower = f.name.lower()
    if lower.endswith(".pdf"):
        return _pdf(f.data)
    if lower.endswith(".docx"):
        return _docx(f.data)
    raw = _decode(f.data)
    if lower.endswith((".html", ".htm")):
        return html_to_text(raw), html_title(raw)
    if lower.endswith((".md", ".markdown", ".mdx")):
        meta, body = frontmatter(raw)
        return clean_mdx(body), meta.get("title")
    return raw, None


def _first_heading(text: str) -> str | None:
    m = re.search(r"^#{1,3}\s+(.+)$", text, re.M)
    return m.group(1).strip() if m else None


class UploadConnector:
    kind = "upload"
    label = "Files"

    def __init__(self, files: list[File], *, today: str | None = None):
        self.files = files
        self.today = today or datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")

    def read(self, *, cursor: str | None, report: Report) -> SourceResult:
        notes: list[str] = []
        files = expand(self.files, notes)
        batch = Batch()
        for i, f in enumerate(files):
            report(step="read_upload", done=i, total=len(files))
            try:
                text, own_title = text_of(f)
            except Exception:
                notes.append(f"{f.name} could not be read.")
                continue
            text = text[:MAX_CHARS].strip()
            if not text:
                notes.append(f"{f.name} has no text Canon can read (a scanned PDF needs OCR first).")
                continue
            stem = PurePosixPath(f.name).stem.replace("_", " ").replace("-", " ").strip()
            title = (own_title or _first_heading(text) or stem or f.name)[:200]
            doc_id = f"upload:{slugify(f.name)[:180] or 'file'}"
            body = text if _first_heading(text) else f"# {title}\n{text}"
            passages = passages_of(doc_id, split_passages(body))
            if not passages:
                continue
            updated = (f.modified or self.today)[:10]
            batch.docs.append(Doc(doc_id, "upload", title, "", updated))  # type: ignore[arg-type]
            batch.passages += passages
        report(step="read_upload", done=len(files), total=len(files))
        return SourceResult(batch=batch, cursor=cursor, notes=notes)
